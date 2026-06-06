#!/usr/bin/env python3
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path("/Users/drew/Downloads/tiktok-n8n-workflow-pack")
REVIEW_DIR = ROOT / ".n8n-local-cache" / "review-context"
OUT_DIR = Path("/Users/drew/Downloads/n8n视频")


def latest_review_context(project_id: str) -> Path:
    files = sorted(
        REVIEW_DIR.glob(f"review_context_{project_id}_*.json"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    files = [p for p in files if not p.name.endswith(".submitted.json")]
    if not files:
        raise SystemExit(f"找不到项目 {project_id} 的 review_context JSON")
    return files[0]


def add_kv(doc: Document, key: str, value) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{key}：").bold = True
    p.add_run(str(value or ""))


def english_video_prompt(panel: dict, total: int) -> str:
    return " ".join(
        [
            "Use the provided storyboard panel image as the exact first frame for this image-to-video clip.",
            "Create a realistic TikTok UGC-style ecommerce product video clip, vertical 9:16, natural handheld smartphone footage, natural light, real-life environment.",
            "Animate only small, natural movement from the visible first frame. Keep the same person identity, face, hairstyle, outfit, body proportion, product shape, product color, material, size, and usage logic.",
            "Keep the product as the visual center. Do not redesign the product. Do not change the background into a different location.",
            "Camera movement should be subtle and stable: small handheld sway, slight push-in, slight tilt, or gentle follow movement only.",
            "Do not add text, captions, subtitles, UI, watermark, logo overlays, stickers, extra fingers, distorted hands, product morphing, face morphing, or sudden scene cuts.",
            f"Stage: {panel.get('stage', '')}.",
            f"Shot order: {panel.get('shot_order', '')} of {total}.",
            f"Local voiceover reference only, do not render on-screen text: {panel.get('optional_voiceover_local', '') or extract_voiceover(panel.get('video_prompt', ''))}.",
            f"Continuity reference: {panel.get('continuity_requirements', '')}.",
        ]
    )


def extract_voiceover(text: str) -> str:
    marker = "Voiceover (local):"
    if marker not in str(text):
        return ""
    return str(text).split(marker, 1)[1].split("Continuity:", 1)[0].strip()


def main() -> None:
    project_id = sys.argv[1] if len(sys.argv) > 1 else ""
    if not project_id:
        raise SystemExit("用法：python export_project_prompt_doc.py proj_xxx")

    context_path = latest_review_context(project_id)
    ctx = json.loads(context_path.read_text("utf-8"))
    panels = ctx.get("panel_review_pack") or []
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading(f"{project_id} 脚本与视频提示词检查文档", 0)
    doc.add_paragraph("用途：给人工检查脚本框架、分镜图、视频模型提示词是否符合产品、目标市场和创作任务。")

    doc.add_heading("一、项目基础信息", level=1)
    add_kv(doc, "项目 ID", ctx.get("project_id"))
    add_kv(doc, "产品", ctx.get("product_name"))
    add_kv(doc, "目标市场", ctx.get("target_market"))
    add_kv(doc, "目标语言", ctx.get("target_language"))
    add_kv(doc, "创作任务类型", ctx.get("creative_task_type"))
    add_kv(doc, "已选择创意方向", ctx.get("selected_concept_name"))
    add_kv(doc, "审核 JSON", context_path)

    doc.add_heading("二、分镜头脚本框架", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["镜头", "阶段", "时长", "场景/动作", "产品状态", "镜头运动", "目标语言口播"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for panel in panels:
        prompt = str(panel.get("video_prompt", ""))
        row = table.add_row().cells
        row[0].text = str(panel.get("shot_order") or panel.get("shot_id") or "")
        row[1].text = str(panel.get("stage") or "")
        row[2].text = str(panel.get("duration") or "")
        row[3].text = prompt
        row[4].text = str(panel.get("product_state") or "")
        row[5].text = str(panel.get("camera_movement") or "")
        row[6].text = extract_voiceover(prompt)

    doc.add_heading("三、逐镜 Video Prompt 检查", level=1)
    doc.add_paragraph("说明：中文脚本用于人工理解；真正建议传给视频模型的 Video Prompt 应使用英文。")
    for panel in panels:
        doc.add_heading(f"镜头 {panel.get('shot_order') or panel.get('shot_id')}", level=2)
        add_kv(doc, "阶段", panel.get("stage"))
        add_kv(doc, "原始脚本/旧版提示词", panel.get("video_prompt"))
        add_kv(doc, "建议英文 Video Prompt", english_video_prompt(panel, len(panels)))
        img_path = panel.get("panel_preview_path") or panel.get("panel_image_path")
        if img_path and Path(img_path).exists():
            try:
                doc.add_picture(img_path, width=Inches(1.4))
            except Exception:
                add_kv(doc, "分镜图路径", img_path)

    out = OUT_DIR / f"{project_id}_脚本与视频提示词检查.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
